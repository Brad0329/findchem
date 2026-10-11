"""MSDS 성분 정리표 — 모델이 뽑은 성분(extract.json)과 classify_cas 응답을 합쳐 표를 만든다.

  python build_table.py cas   extract.json                        → classify_cas에 넣을 CAS 배열(JSON)
  python build_table.py table extract.json classify.json          → 화면용 마크다운 표
  python build_table.py table extract.json classify.json --xlsx out.xlsx   → 위 + 엑셀 파일

extract.json 형식:
  {"products": [{"product": "제품명", "file": "원본.pdf",
                 "components": [{"name": "성분", "cas": "CAS 또는 빈 문자열", "content": "함유량 원문", "note": "선택"}]}]}

분류 문구ㆍ비고ㆍ색칠은 여기서만 만든다 — 모델이 손으로 옮기지 않게(틀린다).
어긋나면 표를 만들지 않고 stderr에 이유를 쓰고 종료 코드 1.
"""

from __future__ import annotations

import datetime
import json
import sys

COLUMNS = ["제품", "구분", "성분", "CAS", "함유량", "물질 분류", "비고", "원본 파일"]

# classify_cas의 classes(고시 용어) → 표에 쓰는 짧은 이름. 순서는 도구가 이미 사고대비물질 먼저로 준다.
CLASS_SHORT = {
    "사고대비물질": "사고대비물질",
    "인체급성유해성": "인체급성",
    "인체만성유해성": "인체만성",
    "생태유해성": "생태",
}
# 별표2 행의 구분 → 함량기준 앞말. 별표3 행은 구분이 빈 문자열이다.
KIND_LABEL = {"": "사고대비", "급성": "급성", "만성": "만성", "생태": "생태"}

NOT_LISTED = "비해당"
NO_CAS = "확인 불가(CAS 없음)"
BAD_CAS = "확인 불가(CAS 형식 오류)"


class TableError(Exception):
    """입력이 어긋나 표를 만들 수 없다."""


def load_json(path: str) -> dict:
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def products_of(extract: dict) -> list[dict]:
    products = extract.get("products")
    if not isinstance(products, list) or not products:
        raise TableError("extract.json의 products가 비어 있다")
    for i, p in enumerate(products, 1):
        comps = p.get("components")
        if not isinstance(comps, list) or not comps:
            raise TableError(f"{i}번째 제품({p.get('product', '?')})에 components가 없다 — 3번 항목을 다시 읽는다")
    return products


def cas_list(extract: dict) -> list[str]:
    """CAS가 있는 성분만, 표 순서대로, 원문 표기 그대로."""
    return [
        c["cas"].strip()
        for p in products_of(extract)
        for c in p["components"]
        if (c.get("cas") or "").strip()
    ]


def content_criteria(entries: list[dict]) -> str:
    """'함량기준 25% 이상' / '함량기준 급성 10% / 만성 0.1% 이상'. 숫자 기준이 없으면 빈 문자열."""
    pairs: list[tuple[str, str]] = []
    for e in entries:
        for r in e.get("rows", []):
            kind, content = r.get("kind", ""), (r.get("content") or "").strip()
            if kind in KIND_LABEL and content and (kind, content) not in pairs:
                pairs.append((kind, content))
    if not pairs:
        return ""
    if len(pairs) == 1:
        return f"함량기준 {pairs[0][1]}% 이상"
    return "함량기준 " + " / ".join(f"{KIND_LABEL[k]} {c}%" for k, c in pairs) + " 이상"


def classify_cell(comp: dict, item: dict | None) -> tuple[str, list[str]]:
    """(물질 분류, 비고 조각들)."""
    note = (comp.get("note") or "").strip()
    if item is None:
        return NO_CAS, ["MSDS에 CAS 미기재"] + ([note] if note else [])

    remarks: list[str] = []
    status = item.get("status")
    if status == "invalid":
        cls = BAD_CAS
        remarks.append(f"CAS 형식 오류: {item.get('reason', '사유 없음')}")
    elif status == "not_listed":
        cls = NOT_LISTED
    elif status == "listed":
        cls = " · ".join(CLASS_SHORT.get(c, c) for c in item.get("classes", []))
        entries = item.get("entries", [])
        crit = content_criteria(entries)
        if crit:
            remarks.append(crit)
        kinds = {r.get("kind") for e in entries for r in e.get("rows", [])}
        if "저확산" in kinds:
            remarks.append("저확산 구분 행 있음(성상 확인 필요)")
        if "용액" in kinds:
            remarks.append("용액 구분 행 있음")
        has_priority = any(e.get("priority") for e in entries)
        if has_priority:
            remarks.append("규정수량은 사고대비물질(별표 3) 우선 적용")
        elif any(e.get("reference_note") for e in entries):
            remarks.append("별표 2 참고값 항목(같은 항목의 다른 CAS가 별표 3에 있음)")
    else:
        raise TableError(f"알 수 없는 status: {status!r} (CAS {item.get('input')})")

    if status != "invalid" and item.get("cas") and item.get("input", "").strip() != item["cas"]:
        remarks.append(f"MSDS 표기 CAS {item['input'].strip()}")
    if item.get("checksum_ok") is False:
        remarks.append("CAS 체크섬 불일치 — 원문 재확인")
    if note:
        remarks.append(note)
    return cls, remarks


def build_rows(extract: dict, classify: dict) -> tuple[list[dict], dict]:
    """성분 행 목록과 요약. 행마다 COLUMNS 키 + 'classified'(제품 색칠용)."""
    products = products_of(extract)
    wanted = cas_list(extract)
    items = classify.get("items")
    if not isinstance(items, list):
        raise TableError("classify.json에 items가 없다 — classify_cas 응답 전체를 그대로 저장한다")
    if len(items) != len(wanted):
        raise TableError(f"CAS 개수가 다르다: 추출 {len(wanted)}개, classify_cas 응답 {len(items)}개")
    for i, (cas, item) in enumerate(zip(wanted, items), 1):
        if (item.get("input") or "").strip() != cas:
            raise TableError(f"{i}번째 CAS가 다르다: 추출 {cas!r}, 응답 input {item.get('input')!r}")

    it = iter(items)
    rows: list[dict] = []
    classified_products = 0
    no_cas = 0
    for p in products:
        comps = p["components"]
        product_rows = []
        for c in comps:
            has_cas = bool((c.get("cas") or "").strip())
            item = next(it) if has_cas else None
            if not has_cas:
                no_cas += 1
            cls, remarks = classify_cell(c, item)
            product_rows.append({
                "제품": "", "구분": "", "원본 파일": "",
                "성분": c.get("name", ""),
                "CAS": (c.get("cas") or "").strip(),
                "함유량": c.get("content", ""),
                "물질 분류": cls,
                "비고": ". ".join(remarks),
                "listed": item is not None and item.get("status") == "listed",
            })
        product_rows[0]["제품"] = p.get("product", "")
        product_rows[0]["구분"] = "혼합물" if len(comps) >= 2 else "단일성분"
        product_rows[0]["원본 파일"] = p.get("file", "")
        classified = any(r["listed"] for r in product_rows)
        classified_products += classified
        for r in product_rows:
            r["classified"] = classified
        rows.extend(product_rows)

    summary = {
        "products": len(products),
        "components": len(rows),
        "classified_products": classified_products,
        "no_cas": no_cas,
        "notice": (classify.get("source") or {}).get("notice", ""),
    }
    return rows, summary


def summary_line(s: dict) -> str:
    return (f"제품 {s['products']}건 · 성분 {s['components']}행 · 분류 해당 제품 {s['classified_products']}건"
            f" · CAS 없는 성분 {s['no_cas']}행")


def footnotes(s: dict, today: str) -> list[str]:
    return [
        "노란색 = 물질 분류에 해당하는 성분이 하나라도 있는 제품.",
        "성분·함유량은 MSDS 3번 항목 원문 표기 그대로다(부등호·범위 포함, 숫자로 바꾸지 않음).",
        f"물질 분류는 {s['notice'] or '「유해화학물질의 규정수량에 관한 규정」'} 별표 2·3을 CAS로 조회한 결과다"
        f"(FindChem classify_cas, {today}).",
        "함유량이 함량기준에 미달하는지는 판정하지 않았다 — 비고의 함량기준과 직접 비교해야 한다.",
        "구분(혼합물/단일성분)은 3번 항목의 성분 수로만 정했다(2개 이상 = 혼합물).",
    ]


def _md(text: str) -> str:
    return str(text).replace("|", "\\|").replace("\n", " ")


def markdown(rows: list[dict], s: dict) -> str:
    out = [summary_line(s), "", "| " + " | ".join(COLUMNS) + " |", "|" + "---|" * len(COLUMNS)]
    for r in rows:
        out.append("| " + " | ".join(_md(r[c]) for c in COLUMNS) + " |")
    return "\n".join(out)


def write_xlsx(rows: list[dict], s: dict, path: str, today: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "MSDS 성분"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for r in rows:
        ws.append([r[c] for c in COLUMNS])
        if r["classified"]:
            for cell in ws[ws.max_row]:
                cell.fill = yellow
    ws.append([])
    for line in footnotes(s, today):
        ws.append([line])
    widths = {"A": 28, "B": 10, "C": 40, "D": 13, "E": 12, "F": 30, "G": 50, "H": 40}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def main(argv: list[str]) -> int:
    try:
        if len(argv) >= 2 and argv[0] == "cas":
            print(json.dumps(cas_list(load_json(argv[1])), ensure_ascii=False))
            return 0
        if len(argv) >= 3 and argv[0] == "table":
            rows, s = build_rows(load_json(argv[1]), load_json(argv[2]))
            today = datetime.date.today().isoformat()
            print(markdown(rows, s))
            if "--xlsx" in argv:
                i = argv.index("--xlsx")
                if i + 1 >= len(argv):
                    raise TableError("--xlsx 뒤에 파일 경로가 없다")
                write_xlsx(rows, s, argv[i + 1], today)
                print(f"\n엑셀 저장: {argv[i + 1]}")
            return 0
    except TableError as e:
        print(f"오류: {e}", file=sys.stderr)
        return 1
    except (OSError, json.JSONDecodeError) as e:
        print(f"오류: 입력 파일을 읽지 못했다 — {e}", file=sys.stderr)
        return 1
    print(__doc__, file=sys.stderr)
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
