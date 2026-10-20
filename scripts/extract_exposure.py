"""A-010 위험노출수준 원천 추출: assets/form6/raw/exposure/ → assets/form6/exposure_tables.json.

원천(받은 날·주소·판은 raw/exposure/sources.json, 조사 결론은 REQUIREMENTS_AGENT A-010):
  aiha_erpg/ERPG.html               AIHA ERPG 표(HTML 원형 그대로 — 라이선스가 원형 배포만 허용). 고지문은 Copyright-and-Licensing.html
  epa_aegl/index.html + <slug>.html EPA AEGL 목록과 물질별 페이지(제목 `CAS (Final|Interim)` 다음의 표, 60 min 열)
  doe_pac/pac_getexcel_20250207.xlsx DOE PAC(Wayback 2025-02-07 — 한국 IP에서 DOE가 막힘). 'Input' 시트, 4행부터 데이터

값은 **원문 문자열 그대로** 둔다(`2,500*`, `0.083`, `NR`). 숫자 해석·고르기는 Dart 생성기(lib/form6/exposure_tables.dart)가 한다.
결정적이다(정렬된 키, 같은 원문 → 같은 바이트). PAC는 대상 CAS(번들 별표 2ㆍ3)만 싣는다 — 3천여 행을 다 실을 이유가 없다.
읽지 못한 것은 조용히 버리지 않고 `skipped`에 사유와 함께 남긴다.

사용: python scripts/extract_exposure.py
"""

from __future__ import annotations

import json
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "assets/form6/raw/exposure"
OUT = ROOT / "assets/form6/exposure_tables.json"
CAS_RE = re.compile(r"\b0*(\d{2,7}-\d\d-\d)\b")


class _Events(HTMLParser):
    """HTML → 문서 순서대로 ('text', str) / ('table', [[셀 글]]) 사건."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.events: list[tuple[str, object]] = []
        self._rows: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell: list[str] | None = None
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        elif tag == "table":
            self._rows = []
        elif tag == "tr" and self._rows is not None:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self._skip = max(0, self._skip - 1)
        elif tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(re.sub(r"\s+", " ", "".join(self._cell)).strip())
            self._cell = None
        elif tag == "tr" and self._row is not None and self._rows is not None:
            self._rows.append(self._row)
            self._row = None
        elif tag == "table" and self._rows is not None:
            self.events.append(("table", self._rows))
            self._rows = None

    def handle_data(self, data):
        if self._skip:
            return
        if self._cell is not None:
            self._cell.append(data)
        else:
            # 표 안이어도 셀 밖 글(caption — AEGL의 `584-84-9 (Final)` 제목이 여기 있다)은 글 사건으로 둔다
            self.events.append(("text", data))


def _events(path: Path) -> list[tuple[str, object]]:
    p = _Events()
    p.feed(path.read_text(encoding="utf-8", errors="strict"))
    return p.events


# AIHA ERPG 표의 CAS 오기 정정 — (표의 물질명, 표의 CAS) → 바른 CAS. 손으로 만든 표라 [check_errata]가 PAC 표로 검증한다.
# 2026-10-20 범위 확장 대조에서 발견: 두 행이 다른 물질의 CAS로 적혀 그 물질에 엉뚱한 ERPG가 붙었다
#   (MDI 101-68-8에 수소화리튬 0.1 mg/m3, 수은 7439-97-6에 무수말레산 5 ppm). PAC 표의 같은 이름 행이 바른 CAS와 같은 ERPG 값을 싣고 있다.
ERPG_ERRATA = {
    ("Lithium hydride", "101-68-8"): "7580-67-8",
    ("Maleic anhydride", "7439-97-6"): "108-31-6",
}


def check_errata(pac_path: Path) -> list[dict]:
    """정정마다 PAC 표 전체에서 같은 이름으로 시작하는 행의 CAS가 바른 CAS인지 본다. 아니면 멈춘다(정정 표가 틀렸다)."""
    ws = openpyxl.load_workbook(pac_path, read_only=True, data_only=True)["Input"]
    rows = [(_s(r[1]), _s(r[2])) for r in ws.iter_rows(min_row=4, values_only=True) if r and r[1]]
    out = []
    for (name, wrong), right in sorted(ERPG_ERRATA.items()):
        hits = [c for n, c in rows if n.lower().startswith(name.lower())]
        if right not in {m for c in hits for m in CAS_RE.findall(c)}:
            raise SystemExit(f"ERPG 정정 검증 실패: {name} — PAC 표 CAS {hits}에 {right}가 없다")
        out.append({"name": name, "aihaCas": wrong, "correctedCas": right, "checkedBy": "DOE PAC 같은 이름 행"})
    return out


def erpg_conflicts(erpg: dict[str, list[dict]]) -> list[dict]:
    """한 CAS에 이름도 ERPG-2 값도 다른 행이 여럿 — 다른 물질의 CAS로 적힌 오기일 수 있다(별칭 행은 값이 같아 걸리지 않는다)."""
    return [
        {"cas": c, "rows": [(r["name"], r["erpg2"]) for r in rows]}
        for c, rows in sorted(erpg.items())
        if len({r["name"] for r in rows}) > 1 and len({r["erpg2"] for r in rows}) > 1
    ]


def extract_erpg(path: Path, skipped: list[str]) -> dict[str, list[dict]]:
    tables = [t for k, t in _events(path) if k == "table"]
    if len(tables) != 1:
        raise SystemExit(f"ERPG: 표가 {len(tables)}개 — 1개여야 한다(페이지 모양이 바뀌었다)")
    rows = tables[0]
    hdr = rows[0]
    for need in ("CAS", "Chemical Name", "ERPG-2", "Last Update"):
        if need not in hdr:
            raise SystemExit(f"ERPG: 머리글에 {need!r} 없음: {hdr}")
    out: dict[str, list[dict]] = {}
    for r in rows[1:]:
        if len(r) != len(hdr):
            skipped.append(f"ERPG 행 칸 수 {len(r)} ≠ {len(hdr)}: {r[:3]}")
            continue
        d = dict(zip(hdr, r))
        cas = sorted(set(CAS_RE.findall(d["CAS"])))
        if not cas:
            skipped.append(f"ERPG CAS 없음: {d['Chemical Name']} ({d['CAS']!r})")
            continue
        fixed = [ERPG_ERRATA.get((d["Chemical Name"], c), c) for c in cas]
        for c in fixed:
            rec = {"name": d["Chemical Name"], "casCell": d["CAS"], "erpg2": d["ERPG-2"], "lastUpdate": d["Last Update"]}
            if fixed != cas:
                rec["casCorrected"] = "AIHA 표의 CAS 오기 정정(ERPG_ERRATA)"
            out.setdefault(c, []).append(rec)
    return out


_HEAD = re.compile(r"(\d{2,7}-\d\d-\d)\s*\(\s*(Final|Interim|Proposed)\s*\)")


# 단위 줄: `ppm`, `ppm 4/20/04`, `(ppm) 11/13/06`, `(mg/m3) 11/30/06`(2026-10-19 받은 69쪽)
_UNIT = re.compile(r"^\(?\s*(ppm|mg/m3|mg/m³|mg/㎥)\s*\)?", re.I)


class _Links(HTMLParser):
    """목록 페이지 → 행마다 (셀 글 이어 붙임, 첫 링크)."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[tuple[str, str | None]] = []
        self._text: list[str] | None = None
        self._href: str | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._text, self._href = [], None
        elif tag == "a" and self._text is not None and self._href is None:
            self._href = dict(attrs).get("href")

    def handle_endtag(self, tag):
        if tag == "tr" and self._text is not None:
            self.rows.append((" ".join(self._text), self._href))
            self._text = None

    def handle_data(self, data):
        if self._text is not None:
            self._text.append(data)


def aegl_index(path: Path) -> dict[str, list[str]]:
    """목록 페이지 → 페이지 이름(slug.html) → 그 행의 CAS들. 제목에 CAS가 없는 물질 페이지(아크릴로니트릴)에 쓴다."""
    p = _Links()
    p.feed(path.read_text(encoding="utf-8"))
    out: dict[str, list[str]] = {}
    for text, href in p.rows:
        if href and href.startswith("/aegl/"):
            out.setdefault(href.rsplit("/", 1)[-1] + ".html", []).extend(CAS_RE.findall(text))
    return out


def extract_aegl_page(path: Path, index: dict[str, list[str]], skipped: list[str]) -> list[dict]:
    """물질별 페이지 → [{cas, status, unit, aegl2_60, footnote, page}]. 값 표(`60 min` 열이 있는 표)마다 하나.

    제목 `CAS (Final|Interim)`은 표 앞 글·표 caption·표 첫 칸 어디에나 온다(2026-10-19 받은 69쪽). 제목이 없는 쪽은
    목록 페이지의 CAS를 쓰고 상태는 `unstated`(아크릴로니트릴).
    """
    out = []
    text_since: list[str] = []
    events = _events(path)
    for i, (kind, val) in enumerate(events):
        if kind == "text":
            text_since.append(val)  # type: ignore[arg-type]
            continue
        rows: list[list[str]] = val  # type: ignore[assignment]
        hdr = next((r for r in rows if "60 min" in r), None)
        heads = _HEAD.findall(re.sub(r"\s+", " ", "".join(text_since) + " " + " ".join(" ".join(r) for r in rows)))
        text_since = []
        if hdr is None:
            continue  # 머리글·꼬리말 표 등 AEGL 값 표가 아니다
        if heads:
            cas, status = heads[-1]
        else:
            listed = sorted(set(index.get(path.name, [])))
            if len(listed) != 1:
                skipped.append(f"AEGL {path.name}: 제목에 CAS가 없고 목록 페이지 CAS도 하나가 아님 {listed}")
                continue
            cas, status = listed[0], "unstated"
        unit_rows = [r[0] for r in rows if r and _UNIT.match(r[0]) and not any(r[1:])]
        unit = _UNIT.match(unit_rows[0]).group(1) if unit_rows else ""  # type: ignore[union-attr]
        a2 = next((r for r in rows if r and re.sub(r"\s+", " ", r[0]) == "AEGL 2"), None)
        if a2 is None or len(a2) != len(hdr):
            skipped.append(f"AEGL {path.name} {cas}: AEGL 2 행 없음·칸 수 다름")
            continue
        foot = []
        for k2, v2 in events[i + 1 :]:
            if k2 != "text":
                break
            foot.append(v2)
        footnote = re.sub(r"\s+", " ", "".join(foot)).strip()  # type: ignore[arg-type]
        out.append(
            {
                "cas": cas,
                "status": status,
                "unit": unit,
                "aegl2_60": a2[hdr.index("60 min")],
                "unitRow": unit_rows[0] if unit_rows else "",
                "footnote": footnote[:300],
                "page": path.name,
            }
        )
    return out


def extract_aegl(dir_: Path, skipped: list[str]) -> dict[str, list[dict]]:
    out: dict[str, list[dict]] = {}
    index = aegl_index(dir_ / "index.html")
    for f in sorted(dir_.glob("*.html")):
        if f.name == "index.html":
            continue
        recs = extract_aegl_page(f, index, skipped)
        if not recs:
            skipped.append(f"AEGL {f.name}: 값 표를 찾지 못함")
        for r in recs:
            out.setdefault(r.pop("cas"), []).append(r)
    return out


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def extract_pac(path: Path, targets: set[str], skipped: list[str]) -> dict[str, list[dict]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Input"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [_s(c) for c in rows[0]]
    if not hdr[2].startswith("CAS") or not hdr[17].startswith("Source of PACs") or hdr[18] != "Units" or _s(rows[1][15]) != "PAC-2":
        raise SystemExit(f"PAC: 열 위치가 바뀌었다: {hdr[:20]}")
    out: dict[str, list[dict]] = {}
    for r in rows[3:]:
        if not r or (r[0] is None and r[1] is None):
            continue
        for c in set(CAS_RE.findall(_s(r[2]))):
            if c in targets:
                out.setdefault(c, []).append(
                    {"no": _s(r[0]), "name": _s(r[1]), "casCell": _s(r[2]), "pac2": _s(r[15]), "source": _s(r[17]), "unit": _s(r[18])}
                )
    return out


def targets_from_bundle() -> set[str]:
    data = json.loads((ROOT / "assets/data/findchem_data.json").read_text(encoding="utf-8"))
    # 별표 2ㆍ3 살아 있는 항목 전부(2026-10-20 범위 확장 — 처음엔 별표 3만)
    return {c for e in data["entries"] if not e["deleted"] for c in e["cas"]}


def aiha_notice(path: Path) -> str:
    """고지문 원문 — 'Copyright & Licensing' 제목 다음부터 'AIHA reserves …' 문장까지(라이선스가 재배포 때 포함을 요구한다)."""
    text = "\n".join(t.strip() for k, t in _events(path) if k == "text" and str(t).strip())
    start = text.find("© 2025 American Industrial Hygiene Association")
    end_key = "AIHA reserves the right to require licensing fees or other remuneration for any commercial use."
    end = text.find(end_key)
    if start < 0 or end < 0:
        raise SystemExit("AIHA 고지문을 찾지 못함 — Copyright-and-Licensing.html 모양이 바뀌었다")
    return text[start : end + len(end_key)]


def build() -> dict:
    skipped: list[str] = []
    targets = targets_from_bundle()
    pac_path = RAW / "doe_pac/pac_getexcel_20250207.xlsx"
    errata = check_errata(pac_path)
    erpg = extract_erpg(RAW / "aiha_erpg/ERPG.html", skipped)
    aegl = extract_aegl(RAW / "epa_aegl", skipped)
    pac = extract_pac(pac_path, targets, skipped)
    return {
        "generatedBy": "scripts/extract_exposure.py",
        "sources": json.loads((RAW / "sources.json").read_text(encoding="utf-8")),
        "aihaNotice": aiha_notice(RAW / "aiha_erpg/Copyright-and-Licensing.html"),
        "erpgErrata": errata,
        "erpgConflicts": erpg_conflicts(erpg),
        "erpg": dict(sorted(erpg.items())),
        "aegl": dict(sorted(aegl.items())),
        "pac": dict(sorted(pac.items())),
        "skipped": skipped,
    }


def render(d: dict) -> str:
    return json.dumps(d, ensure_ascii=False, indent=1) + "\n"


def main() -> int:
    d = build()
    OUT.write_text(render(d), encoding="utf-8", newline="\n")
    sys.stdout.reconfigure(encoding="utf-8")
    print(f"exposure_tables.json: ERPG CAS {len(d['erpg'])}, AEGL CAS {len(d['aegl'])}, PAC CAS {len(d['pac'])}, 건너뜀 {len(d['skipped'])}, "
          f"ERPG 정정 {len(d['erpgErrata'])}, ERPG 충돌 {len(d['erpgConflicts'])}")
    for c in d["erpgConflicts"]:
        print("  ERPG 충돌(오기 의심):", c)
    for s in d["skipped"]:
        print("  건너뜀:", s)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
