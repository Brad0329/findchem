"""F-007 GHS 그림문자: assets/유해성 분류.xlsx(사용자 제공 2026-09-15) → 그림 9장 + Dart 대응표.

엑셀 구조(실측): 3행에 코드(GHS01~GHS09, 열마다 하나), 4행에 유해성 분류(줄바꿈으로 여러 개), 2행에 그림.
그림은 셀에 묶이지 않은 절대 위치(absoluteAnchor)이고 **파일 번호 순서가 열 순서와 다르다**(image4 = GHS05,
image5 = GHS04) — 그래서 x 좌표 순으로 열에 짝짓는다. 짝은 그림을 눈으로 대조해 확인했다(2026-09-15).

출력(둘 다 커밋한다 — 앱은 엑셀을 읽지 않는다):
  assets/data/ghs/GHSxx.png        160px로 줄인 그림(웹 번들 크기 — 원본은 1000px·장당 약 100KB)
  lib/lookup/ghs_pictogram_data.dart  코드 → 유해성 분류 목록 (생성물 — 손으로 고치지 않는다)

실행: python scripts/extract_ghs_pictograms.py
"""
import io
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "assets" / "유해성 분류.xlsx"
OUT_DIR = ROOT / "assets" / "data" / "ghs"
DART_OUT = ROOT / "lib" / "lookup" / "ghs_pictogram_data.dart"
SIZE = 160  # 화면 72px × 고해상도 2배 여유. 256px은 장당 약 30KB였다
CODE = re.compile(r"^GHS\d{2}$")


def read_codes_and_labels(path: Path) -> list[tuple[str, list[str]]]:
    ws = openpyxl.load_workbook(path).worksheets[0]
    for row in ws.iter_rows():
        codes = [c for c in row if isinstance(c.value, str) and CODE.match(c.value.strip())]
        if not codes:
            continue
        out = []
        for c in codes:
            label = ws.cell(row=c.row + 1, column=c.column).value or ""
            lines = [s.strip() for s in str(label).split("\n") if s.strip()]
            if not lines:
                raise SystemExit(f"{c.value}: 유해성 분류가 비어 있다(셀 {c.coordinate} 아래)")
            out.append((c.value.strip(), lines))
        return out
    raise SystemExit("GHS 코드 행을 찾지 못했다")


def read_images_by_x(path: Path) -> list[bytes]:
    z = zipfile.ZipFile(path)
    drawing = z.read("xl/drawings/drawing1.xml").decode("utf-8")
    rels = dict(
        re.findall(
            r'Id="(rId\d+)"[^>]*Target="\.\./media/([^"]+)"',
            z.read("xl/drawings/_rels/drawing1.xml.rels").decode("utf-8"),
        )
    )
    anchors = re.findall(
        r'<xdr:absoluteAnchor><xdr:pos x="(\d+)" y="(\d+)"/>.*?r:embed="(rId\d+)"', drawing, flags=re.S
    )
    if len({y for _, y, _ in anchors}) != 1:
        raise SystemExit(f"그림이 한 줄에 있지 않다: {anchors}")
    anchors.sort(key=lambda a: int(a[0]))
    return [z.read(f"xl/media/{rels[rid]}") for _, _, rid in anchors]


def main() -> None:
    entries = read_codes_and_labels(SOURCE)
    images = read_images_by_x(SOURCE)
    if len(entries) != len(images):
        raise SystemExit(f"코드 {len(entries)}개와 그림 {len(images)}장이 맞지 않는다")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for (code, _), raw in zip(entries, images):
        img = Image.open(io.BytesIO(raw)).convert("RGBA").resize((SIZE, SIZE), Image.LANCZOS)
        img.save(OUT_DIR / f"{code}.png", optimize=True)

    body = "\n".join(
        f"  '{code}': [{', '.join(repr(s).replace(chr(34), chr(39)) for s in lines)}],"
        for code, lines in entries
    )
    DART_OUT.write_text(
        "// 생성물: scripts/extract_ghs_pictograms.py (원본 assets/유해성 분류.xlsx) — 손으로 고치지 않는다.\n"
        "\n"
        "/// GHS 그림문자 코드 → 유해성 분류(엑셀 한 칸의 줄들). 그림은 [ghsPictogramAsset].\n"
        "const ghsPictogramLabels = <String, List<String>>{\n"
        f"{body}\n"
        "};\n"
        "\n"
        "String ghsPictogramAsset(String code) => 'assets/data/ghs/$code.png';\n",
        encoding="utf-8",
        newline="\n",
    )
    sys.stdout.reconfigure(encoding="utf-8")
    for code, lines in entries:
        size = (OUT_DIR / f"{code}.png").stat().st_size
        print(code, lines, f"{size} bytes")


if __name__ == "__main__":
    main()
