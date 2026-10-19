"""A-009 MSDS 성분 정리표 스크립트 — 정답지 10msds.xlsx(Drive) 10건의 추출 결과 + 실제 classify_cas 응답으로 돈다."""

import copy
import json
import os
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "agent_skills" / "msds-composition" / "scripts" / "build_table.py"
FIX = ROOT / "tests" / "fixtures" / "msds_skill"
sys.path.insert(0, str(SCRIPT.parent))

import build_table as bt  # noqa: E402

# 정답지 10msds.xlsx의 물질 분류 열(17행, 2026-10-11 사용자 확인본)
EXPECTED_CLASS = [
    "비해당", "비해당", "인체급성",
    "인체만성", "비해당", "비해당",
    "확인 불가(CAS 없음)", "인체급성",
    "인체급성",
    "인체만성",
    "인체급성", "확인 불가(CAS 없음)",
    "인체급성 · 인체만성", "비해당",
    "사고대비물질 · 인체급성 · 인체만성",
    "비해당",
    "인체급성",
]


@pytest.fixture
def extract():
    return json.loads((FIX / "extract_10msds.json").read_text(encoding="utf-8"))


@pytest.fixture
def classify():
    return json.loads((FIX / "classify_10msds.json").read_text(encoding="utf-8"))


def by_name(rows, name):
    return next(r for r in rows if r["성분"] == name)


def test_cas_list_keeps_order_and_raw_form(extract):
    cas = bt.cas_list(extract)
    assert len(cas) == 15
    assert cas[0] == "123-86-4" and cas[-1] == "2855-13-2"
    assert "000065-85-0" in cas


def test_rows_and_columns(extract, classify):
    rows, s = bt.build_rows(extract, classify)
    assert len(rows) == 17
    assert [r["제품"] for r in rows[:3]] == ["AT-2150", "", ""]
    assert rows[0]["원본 파일"] == "AT-2150(K)MSDS.pdf" and rows[1]["원본 파일"] == ""
    assert s == {"products": 10, "components": 17, "classified_products": 9, "no_cas": 2,
                 "notice": classify["source"]["notice"]}


def test_class_column_matches_answer_key(extract, classify):
    rows, _ = bt.build_rows(extract, classify)
    assert [r["물질 분류"] for r in rows] == EXPECTED_CLASS


def test_remarks(extract, classify):
    rows, _ = bt.build_rows(extract, classify)
    assert by_name(rows, "아크릴로니트릴(Acrylonitrile)")["비고"] == (
        "함량기준 사고대비 0.1% / 급성 1% / 만성 0.1% 이상. 규정수량은 사고대비물질(별표 3) 우선 적용")
    assert by_name(rows, "MEKO")["비고"] == "함량기준 급성 10% / 만성 0.1% 이상"
    assert by_name(rows, "벤조산")["비고"] == (
        "함량기준 10% 이상. 저확산 구분 행 있음(성상 확인 필요). MSDS 표기 CAS 000065-85-0")
    assert by_name(rows, "톨루엔 디이소시아네이트")["비고"] == (
        "함량기준 1% 이상. 별표 2 참고값 항목(같은 항목의 다른 CAS가 별표 3에 있음)")
    assert by_name(rows, "영업비밀")["비고"] == "MSDS에 CAS 미기재. CAS 칸 표기 '- / -'"
    assert by_name(rows, "부틸 아세테이트")["비고"] == ""


def test_kind_mixture_vs_single(extract, classify):
    rows, _ = bt.build_rows(extract, classify)
    kinds = {r["제품"]: r["구분"] for r in rows if r["제품"]}
    assert kinds["염화벤죠일(Benezoyl Chloride)"] == "혼합물"
    assert kinds["시클로헥사논"] == "단일성분"
    assert list(kinds.values()).count("혼합물") == 5


def test_mismatched_response_is_refused(extract, classify):
    short = copy.deepcopy(classify)
    short["items"].pop()
    with pytest.raises(bt.TableError, match="개수"):
        bt.build_rows(extract, short)
    swapped = copy.deepcopy(classify)
    swapped["items"][0], swapped["items"][1] = swapped["items"][1], swapped["items"][0]
    with pytest.raises(bt.TableError, match="1번째"):
        bt.build_rows(extract, swapped)


def test_invalid_and_checksum(extract, classify):
    ex = copy.deepcopy(extract)
    ex["products"][0]["components"][0]["cas"] = "203-466-5"
    cl = copy.deepcopy(classify)
    cl["items"][0] = {"input": "203-466-5", "status": "invalid", "reason": "CAS 형식이 아닙니다"}
    cl["items"][1]["checksum_ok"] = False
    rows, _ = bt.build_rows(ex, cl)
    assert rows[0]["물질 분류"] == "확인 불가(CAS 형식 오류)"
    assert "CAS 형식이 아닙니다" in rows[0]["비고"]
    assert "체크섬 불일치" in rows[1]["비고"]


def test_markdown_summary_first(extract, classify):
    rows, s = bt.build_rows(extract, classify)
    md = bt.markdown(rows, s).splitlines()
    assert md[0] == "제품 10건 · 성분 17행 · 분류 해당 제품 9건 · CAS 없는 성분 2행"
    assert len([l for l in md if l.startswith("| ") and "---" not in l]) == 18  # 머리글 + 17행


def test_xlsx(extract, classify, tmp_path):
    from openpyxl import load_workbook

    out = tmp_path / "out.xlsx"
    rows, s = bt.build_rows(extract, classify)
    bt.write_xlsx(rows, s, str(out), "2026-10-11")
    ws = load_workbook(out).active
    assert [c.value for c in ws[1]] == bt.COLUMNS
    assert ws.cell(2, 6).value == "비해당"
    yellow = [r for r in range(2, 19) if ws.cell(r, 1).fill.fgColor.rgb.endswith("FFF2CC")]
    assert len(yellow) == 16  # 분류 해당 9제품의 행 전부(TEA 1행만 빠진다)
    assert not ws.cell(17, 1).fill.fgColor.rgb.endswith("FFF2CC")  # TEA
    assert ws.cell(20, 1).value.startswith("노란색")
    assert "2026-10-11" in ws.cell(22, 1).value


def test_cli_error_exit_code(tmp_path):
    bad = tmp_path / "c.json"
    bad.write_text(json.dumps({"items": []}), encoding="utf-8")
    # Windows에서는 자식의 stderr 기본 인코딩이 cp949라 utf-8로 읽으면 깨진다 — 자식에게 utf-8로 쓰게 한다
    r = subprocess.run([sys.executable, str(SCRIPT), "table", str(FIX / "extract_10msds.json"), str(bad)],
                       capture_output=True, text=True, encoding="utf-8", env={**os.environ, "PYTHONIOENCODING": "utf-8"})
    assert r.returncode == 1 and "개수" in r.stderr
